#!/usr/bin/env python3
"""
Скрипт для экспорта информации о таблицах Snowflake в Excel файл
Создает подробную документацию структуры базы данных
"""

import os
import sys
from datetime import datetime
from dotenv import load_dotenv
import snowflake.connector
from snowflake.connector.errors import ProgrammingError

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    print("Ошибка: OpenPyXL не установлен. Установите: pip install openpyxl")
    sys.exit(1)

# Загружаем переменные окружения
load_dotenv()

# Конфигурация Snowflake
SNOWFLAKE_CONFIG = {
    'user': os.getenv('SNOWFLAKE_USER'),
    'password': os.getenv('SNOWFLAKE_PASSWORD'),
    'account': os.getenv('SNOWFLAKE_ACCOUNT'),
    'warehouse': os.getenv('SNOWFLAKE_WAREHOUSE', 'COMPUTE_WH'),
    'database': os.getenv('SNOWFLAKE_DATABASE', 'SPOTIFY_DB'),
    'schema': os.getenv('SNOWFLAKE_SCHEMA', 'PUBLIC')
}


def get_table_list(conn, cur):
    """Получает список всех таблиц в схеме"""
    try:
        cur.execute("""
            SELECT TABLE_NAME, TABLE_TYPE
            FROM INFORMATION_SCHEMA.TABLES
            WHERE TABLE_SCHEMA = %s
            AND TABLE_TYPE = 'BASE TABLE'
            ORDER BY TABLE_NAME
        """, (SNOWFLAKE_CONFIG['schema'],))
        return cur.fetchall()
    except Exception as e:
        print(f"Ошибка при получении списка таблиц: {e}")
        return []


def get_table_columns(conn, cur, table_name):
    """Получает информацию о колонках таблицы"""
    try:
        cur.execute("""
            SELECT 
                COLUMN_NAME,
                DATA_TYPE,
                CHARACTER_MAXIMUM_LENGTH,
                NUMERIC_PRECISION,
                NUMERIC_SCALE,
                IS_NULLABLE,
                COLUMN_DEFAULT,
                COLUMN_COMMENT
            FROM INFORMATION_SCHEMA.COLUMNS
            WHERE TABLE_SCHEMA = %s
            AND TABLE_NAME = %s
            ORDER BY ORDINAL_POSITION
        """, (SNOWFLAKE_CONFIG['schema'], table_name))
        return cur.fetchall()
    except Exception as e:
        print(f"Ошибка при получении колонок для таблицы {table_name}: {e}")
        return []


def get_table_constraints(conn, cur, table_name):
    """Получает информацию об ограничениях таблицы"""
    try:
        cur.execute("""
            SELECT 
                CONSTRAINT_NAME,
                CONSTRAINT_TYPE
            FROM INFORMATION_SCHEMA.TABLE_CONSTRAINTS
            WHERE TABLE_SCHEMA = %s
            AND TABLE_NAME = %s
        """, (SNOWFLAKE_CONFIG['schema'], table_name))
        return cur.fetchall()
    except Exception as e:
        print(f"Ошибка при получении ограничений для таблицы {table_name}: {e}")
        return []


def get_primary_keys(conn, cur, table_name):
    """Получает информацию о первичных ключах"""
    try:
        cur.execute("""
            SELECT 
                kcu.COLUMN_NAME,
                kcu.ORDINAL_POSITION
            FROM INFORMATION_SCHEMA.TABLE_CONSTRAINTS tc
            JOIN INFORMATION_SCHEMA.KEY_COLUMN_USAGE kcu
                ON tc.CONSTRAINT_NAME = kcu.CONSTRAINT_NAME
                AND tc.TABLE_SCHEMA = kcu.TABLE_SCHEMA
                AND tc.TABLE_NAME = kcu.TABLE_NAME
            WHERE tc.TABLE_SCHEMA = %s
            AND tc.TABLE_NAME = %s
            AND tc.CONSTRAINT_TYPE = 'PRIMARY KEY'
            ORDER BY kcu.ORDINAL_POSITION
        """, (SNOWFLAKE_CONFIG['schema'], table_name))
        return cur.fetchall()
    except Exception as e:
        return []


def get_row_count(conn, cur, table_name):
    """Получает количество строк в таблице"""
    try:
        cur.execute(f"SELECT COUNT(*) FROM {table_name}")
        result = cur.fetchone()
        return result[0] if result else 0
    except Exception as e:
        return None


def format_data_type(data_type, max_length, precision, scale):
    """Форматирует тип данных для отображения"""
    if data_type in ('VARCHAR', 'CHAR', 'STRING', 'TEXT'):
        if max_length:
            return f"{data_type}({max_length})"
        return data_type
    elif data_type in ('NUMBER', 'DECIMAL', 'NUMERIC'):
        if precision and scale:
            return f"{data_type}({precision},{scale})"
        elif precision:
            return f"{data_type}({precision})"
        return data_type
    elif data_type == 'ARRAY':
        return "ARRAY"
    else:
        return data_type


def create_excel_export():
    """Создает Excel файл с информацией о таблицах Snowflake"""
    print("=" * 60)
    print("Экспорт структуры таблиц Snowflake в Excel")
    print("=" * 60)
    print()
    
    if not OPENPYXL_AVAILABLE:
        print("✗ OpenPyXL не установлен")
        return False
    
    # Подключаемся к Snowflake
    try:
        print("Подключение к Snowflake...")
        conn = snowflake.connector.connect(**SNOWFLAKE_CONFIG)
        cur = conn.cursor()
        cur.execute(f"USE SCHEMA {SNOWFLAKE_CONFIG['schema']}")
        print(f"✓ Подключение установлено (Database: {SNOWFLAKE_CONFIG['database']}, Schema: {SNOWFLAKE_CONFIG['schema']})")
        print()
    except Exception as e:
        print(f"✗ Ошибка подключения к Snowflake: {e}")
        return False
    
    try:
        # Создаем рабочую книгу
        wb = Workbook()
        wb.remove(wb.active)  # Удаляем дефолтный лист
        
        # Стили
        header_fill = PatternFill(start_color="1DB954", end_color="1DB954", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        title_font = Font(bold=True, size=14, color="1DB954")
        
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Получаем список таблиц
        print("Получение списка таблиц...")
        tables = get_table_list(conn, cur)
        print(f"✓ Найдено {len(tables)} таблиц")
        print()
        
        if not tables:
            print("Таблицы не найдены!")
            return False
        
        # Создаем лист со списком таблиц
        ws_summary = wb.create_sheet("Список таблиц", 0)
        ws_summary.append(["Название таблицы", "Тип", "Количество строк", "Количество колонок"])
        
        # Применяем стили к заголовкам
        for cell in ws_summary[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border
        
        # Обрабатываем каждую таблицу
        print("Обработка таблиц...")
        print("-" * 60)
        
        for idx, (table_name, table_type) in enumerate(tables, 1):
            print(f"[{idx}/{len(tables)}] Обработка таблицы: {table_name}")
            
            # Получаем информацию о таблице
            columns = get_table_columns(conn, cur, table_name)
            constraints = get_table_constraints(conn, cur, table_name)
            primary_keys = get_primary_keys(conn, cur, table_name)
            row_count = get_row_count(conn, cur, table_name)
            
            # Добавляем в сводный лист
            ws_summary.append([
                table_name,
                table_type,
                row_count if row_count is not None else "N/A",
                len(columns)
            ])
            
            # Создаем лист для таблицы
            ws = wb.create_sheet(table_name)
            
            # Заголовок таблицы
            ws.merge_cells('A1:H1')
            title_cell = ws['A1']
            title_cell.value = f"Таблица: {table_name}"
            title_cell.font = title_font
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Информация о таблице
            ws.append([])
            ws.append(["Информация о таблице"])
            ws.append(["Тип:", table_type])
            ws.append(["Количество строк:", row_count if row_count is not None else "N/A"])
            ws.append(["Количество колонок:", len(columns)])
            
            pk_columns = [pk[0] for pk in primary_keys]
            if pk_columns:
                ws.append(["Первичный ключ:", ", ".join(pk_columns)])
            
            constraint_types = {}
            for constraint_name, constraint_type in constraints:
                if constraint_type not in constraint_types:
                    constraint_types[constraint_type] = []
                constraint_types[constraint_type].append(constraint_name)
            
            if constraint_types:
                ws.append(["Ограничения:"])
                for constraint_type, names in constraint_types.items():
                    ws.append([f"  {constraint_type}:", ", ".join(names)])
            
            ws.append([])
            ws.append([])
            
            # Заголовки колонок
            headers = [
                "Порядковый номер",
                "Название колонки",
                "Тип данных",
                "Макс. длина",
                "Точность/Масштаб",
                "NULL разрешен",
                "Значение по умолчанию",
                "Комментарий"
            ]
            ws.append(headers)
            
            # Применяем стили к заголовкам
            for cell in ws[len(ws['A'])]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                cell.border = border
            
            # Данные о колонках
            for col_idx, col in enumerate(columns, 1):
                col_name = col[0]
                data_type = col[1]
                max_length = col[2]
                precision = col[3]
                scale = col[4]
                is_nullable = col[5]
                default_value = col[6]
                comment = col[7]
                
                formatted_type = format_data_type(data_type, max_length, precision, scale)
                
                row = [
                    col_idx,
                    col_name,
                    formatted_type,
                    max_length if max_length else "",
                    f"{precision},{scale}" if precision and scale else (precision if precision else ""),
                    "Да" if is_nullable == "YES" else "Нет",
                    str(default_value) if default_value else "",
                    comment if comment else ""
                ]
                
                ws.append(row)
                
                # Применяем стили к ячейкам
                for cell in ws[len(ws['A'])]:
                    cell.border = border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            
            # Автоподбор ширины колонок
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if cell.value:
                            cell_length = len(str(cell.value))
                            if cell_length > max_length:
                                max_length = cell_length
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Автоподбор ширины для сводного листа
        for column in ws_summary.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws_summary.column_dimensions[column_letter].width = adjusted_width
        
        # Применяем стили к сводному листу
        for row in ws_summary.iter_rows(min_row=2, max_row=ws_summary.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center")
        
        # Сохраняем файл
        filename = f"snowflake_schema_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(filename)
        
        print()
        print("=" * 60)
        print(f"✓ Файл успешно создан: {filename}")
        print(f"  - Обработано таблиц: {len(tables)}")
        print(f"  - Создано листов: {len(tables) + 1}")
        print("=" * 60)
        
        return True
        
    except Exception as e:
        print(f"✗ Ошибка при создании Excel файла: {e}")
        import traceback
        traceback.print_exc()
        return False
    finally:
        cur.close()
        conn.close()
        print("\n✓ Соединение с Snowflake закрыто")


if __name__ == "__main__":
    try:
        success = create_excel_export()
        if not success:
            sys.exit(1)
    except KeyboardInterrupt:
        print("\n\n⚠ Прервано пользователем")
        sys.exit(1)
    except Exception as e:
        print(f"\n✗ Неожиданная ошибка: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
