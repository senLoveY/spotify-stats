import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_schema_excel():
    wb = openpyxl.Workbook()
    # Удаляем стандартный лист
    wb.remove(wb.active)

    # Стили
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1DB954", end_color="1DB954", fill_type="solid") # Spotify Green
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    # Данные схемы (Таблица: [Список строк])
    schema_data = {
        "users": [
            ["id", "SERIAL", "PK, Auto-increment", "Уникальный ID пользователя"],
            ["username", "VARCHAR(100)", "UNIQUE, NOT NULL", "Логин"],
            ["password_hash", "VARCHAR(255)", "NULLABLE", "Хэш пароля (Werkzeug)"],
            ["email", "VARCHAR(255)", "UNIQUE", "Email адрес"],
            ["spotify_id", "VARCHAR(100)", "UNIQUE", "ID Spotify аккаунта"],
            ["display_name", "VARCHAR(150)", "", "Отображаемое имя"],
            ["role", "VARCHAR(20)", "CHECK ('user', 'admin')", "Роль пользователя"],
            ["is_blocked", "BOOLEAN", "DEFAULT FALSE", "Статус блокировки"],
            ["email_verified", "BOOLEAN", "DEFAULT FALSE", "Подтвержден ли email"],
            ["created_at", "TIMESTAMP", "DEFAULT NOW()", "Дата регистрации"]
        ],
        "user_sessions": [
            ["id", "SERIAL", "PK", "ID сессии"],
            ["user_id", "INTEGER", "FK (users)", "ID пользователя"],
            ["session_token", "VARCHAR(255)", "UNIQUE, NOT NULL", "Токен сессии"],
            ["ip_address", "VARCHAR(45)", "", "IP адрес клиента"],
            ["user_agent", "TEXT", "", "Данные о браузере"],
            ["created_at", "TIMESTAMP", "DEFAULT NOW()", "Время входа"],
            ["expires_at", "TIMESTAMP", "NOT NULL", "Время истечения"],
            ["is_active", "BOOLEAN", "DEFAULT TRUE", "Активна ли сессия"]
        ],
        "verification_codes": [
            ["id", "SERIAL", "PK", "ID записи"],
            ["user_id", "INTEGER", "FK (users)", "ID пользователя"],
            ["email", "VARCHAR(255)", "NOT NULL", "Email получателя"],
            ["code", "VARCHAR(6)", "LENGTH=6", "Код подтверждения"],
            ["purpose", "VARCHAR(50)", "registration, login...", "Назначение кода"],
            ["expires_at", "TIMESTAMP", "NOT NULL", "Срок действия"],
            ["used", "BOOLEAN", "DEFAULT FALSE", "Использован ли код"]
        ],
        "user_preferences": [
            ["id", "SERIAL", "PK", "ID настройки"],
            ["user_id", "INTEGER", "FK (users), UNIQUE", "ID пользователя"],
            ["language", "VARCHAR(10)", "CHECK ('ru', 'en')", "Язык интерфейса"],
            ["notifications_enabled", "BOOLEAN", "DEFAULT TRUE", "Уведомления вкл/выкл"],
            ["email_notifications", "BOOLEAN", "DEFAULT TRUE", "Email рассылки"],
            ["default_view", "VARCHAR(50)", "dashboard/profile...", "Стартовая страница"],
            ["timezone", "VARCHAR(50)", "DEFAULT UTC", "Часовой пояс"]
        ],
        "feedback": [
            ["id", "SERIAL", "PK", "ID тикета"],
            ["user_id", "INTEGER", "FK (users)", "Автор (может быть NULL)"],
            ["feedback_type", "VARCHAR(50)", "bug, feature...", "Тип обращения"],
            ["title", "VARCHAR(255)", "NOT NULL", "Заголовок"],
            ["message", "TEXT", "NOT NULL", "Текст сообщения"],
            ["status", "VARCHAR(20)", "new, resolved...", "Статус тикета"],
            ["priority", "VARCHAR(10)", "low, high...", "Приоритет"],
            ["created_at", "TIMESTAMP", "DEFAULT NOW()", "Дата создания"]
        ],
        "audit_logs": [
            ["id", "SERIAL", "PK", "ID лога"],
            ["user_id", "INTEGER", "FK (users)", "Кто совершил действие"],
            ["action", "VARCHAR(100)", "NOT NULL", "Код действия"],
            ["details", "TEXT", "", "Детали события"],
            ["timestamp", "TIMESTAMP", "DEFAULT NOW()", "Время события"]
        ],
        "users_history": [
            ["id", "SERIAL", "PK", "ID архива"],
            ["user_id", "INTEGER", "NOT NULL", "ID пользователя"],
            ["operation_type", "VARCHAR(10)", "INSERT/UPDATE/DELETE", "Тип изменения"],
            ["old_value", "JSONB", "", "Данные ДО"],
            ["new_value", "JSONB", "", "Данные ПОСЛЕ"],
            ["operation_date", "TIMESTAMP", "DEFAULT NOW()", "Время изменения"]
        ],
        "feedback_history": [
            ["id", "SERIAL", "PK", "ID архива"],
            ["feedback_id", "INTEGER", "NOT NULL", "ID тикета"],
            ["operation_type", "VARCHAR(10)", "INSERT/UPDATE/DELETE", "Тип изменения"],
            ["old_value", "JSONB", "", "Данные ДО"],
            ["new_value", "JSONB", "", "Данные ПОСЛЕ"],
            ["operation_date", "TIMESTAMP", "DEFAULT NOW()", "Время изменения"]
        ]
    }

    headers = ["Поле", "Тип данных", "Ограничения", "Описание"]

    for table_name, rows in schema_data.items():
        ws = wb.create_sheet(title=table_name[:31]) # Excel ограничивает длину имени листа

        # Заголовок листа (название таблицы)
        ws.merge_cells('A1:D1')
        title_cell = ws['A1']
        title_cell.value = f"Таблица: {table_name}"
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = center_align

        # Заголовки колонок
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Данные
        for row_num, row_data in enumerate(rows, 3):
            for col_num, cell_value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = cell_value
                cell.alignment = left_align
                cell.border = thin_border

        # Настройка ширины колонок
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 40

    filename = "database_schema_full.xlsx"
    wb.save(filename)
    print(f"Файл '{filename}' успешно создан!")

if __name__ == "__main__":
    create_schema_excel()